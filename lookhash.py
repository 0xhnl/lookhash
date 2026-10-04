#!/usr/bin/env python3
"""
lookhash.py - one-shot hash dump -> ntlm.pw lookup -> Excel report

Usage:
    python3 lookhash.py ../all-dump.txt -o results.xlsx
    python3 lookhash.py ../all-dump.txt -o results.xlsx -t nt --no-lookup

Does, in order:
    1. Extract unique NT hashes from a secretsdump-style file
    2. Look them up against the ntlm.pw API (300/request, rate-limit aware)
       and/or test one or more candidate passwords locally (-p)
    3. Write a 2-sheet Excel report (All_Hashes + Cracked_Passwords)

Examples:
    python3 lookhash.py dump.txt -o report.xlsx
    python3 lookhash.py dump.txt -o report.xlsx -p 'P@ssw0rd123' --no-lookup
    python3 lookhash.py dump.txt -o report.xlsx -p 'Summer2024!' -p 'Welcome1'
"""

import os
import re
import sys
import time
import hashlib
import argparse

import requests
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

API_BASE_URL = "https://ntlm.pw/api/lookup"


# --------------------------------------------------------------------------- #
# 1. Extraction                                                               #
# --------------------------------------------------------------------------- #
def extract_nt_hashes(input_file):
    """Return a sorted list of unique NT hashes found in the dump."""
    print(f"[*] Extracting NT hashes from {input_file}...")
    nt_hashes = set()
    pattern = r'(?:[^\s:]+\\)?[^\s:]+:\d+:[a-f0-9]{32}:([a-f0-9]{32}):::'

    try:
        with open(input_file, 'r', encoding='utf-8', errors='ignore') as f:
            for line in f:
                line = line.strip()
                if not line:
                    continue
                match = re.search(pattern, line)
                if match:
                    nt_hashes.add(match.group(1))
                elif re.match(r'^[a-f0-9]{32}$', line):
                    nt_hashes.add(line)
    except FileNotFoundError:
        print(f"[!] Error: File '{input_file}' not found")
        sys.exit(1)

    hashes = sorted(nt_hashes)
    print(f"[+] Extracted {len(hashes)} unique NT hashes")
    return hashes


# --------------------------------------------------------------------------- #
# 2. Lookup                                                                    #
# --------------------------------------------------------------------------- #
def _md4_pure(data: bytes) -> str:
    """Pure-Python MD4 fallback for builds where OpenSSL has dropped md4."""
    def lrot(x, n):
        x &= 0xFFFFFFFF
        return ((x << n) | (x >> (32 - n))) & 0xFFFFFFFF

    a, b, c, d = 0x67452301, 0xEFCDAB89, 0x98BADCFE, 0x10325476
    msg = bytearray(data)
    orig_bits = (len(msg) * 8) & 0xFFFFFFFFFFFFFFFF
    msg.append(0x80)
    while len(msg) % 64 != 56:
        msg.append(0)
    msg += orig_bits.to_bytes(8, 'little')

    for off in range(0, len(msg), 64):
        X = [int.from_bytes(msg[off + 4 * i:off + 4 * i + 4], 'little') for i in range(16)]
        aa, bb, cc, dd = a, b, c, d

        def F(x, y, z): return (x & y) | (~x & z)
        def G(x, y, z): return (x & y) | (x & z) | (y & z)
        def H(x, y, z): return x ^ y ^ z

        for i in range(0, 16, 4):
            a = lrot(a + F(b, c, d) + X[i], 3)
            d = lrot(d + F(a, b, c) + X[i + 1], 7)
            c = lrot(c + F(d, a, b) + X[i + 2], 11)
            b = lrot(b + F(c, d, a) + X[i + 3], 19)
        for i in range(4):
            a = lrot(a + G(b, c, d) + X[i] + 0x5A827999, 3)
            d = lrot(d + G(a, b, c) + X[i + 4] + 0x5A827999, 5)
            c = lrot(c + G(d, a, b) + X[i + 8] + 0x5A827999, 9)
            b = lrot(b + G(c, d, a) + X[i + 12] + 0x5A827999, 13)
        for i in [0, 2, 1, 3]:
            a = lrot(a + H(b, c, d) + X[i] + 0x6ED9EBA1, 3)
            d = lrot(d + H(a, b, c) + X[i + 8] + 0x6ED9EBA1, 9)
            c = lrot(c + H(d, a, b) + X[i + 4] + 0x6ED9EBA1, 11)
            b = lrot(b + H(c, d, a) + X[i + 12] + 0x6ED9EBA1, 15)

        a = (a + aa) & 0xFFFFFFFF
        b = (b + bb) & 0xFFFFFFFF
        c = (c + cc) & 0xFFFFFFFF
        d = (d + dd) & 0xFFFFFFFF

    return b''.join(x.to_bytes(4, 'little') for x in (a, b, c, d)).hex()


def generate_ntlm_hash(password):
    """Return the lowercase NTLM (MD4 of UTF-16LE) hash of a password."""
    data = password.encode('utf-16le')
    try:
        return hashlib.new('md4', data).hexdigest().lower()
    except Exception:
        return _md4_pure(data).lower()


def custom_lookup(hashes, passwords):
    """Test candidate passwords locally. Returns {hash: password} for matches."""
    cracked = {}
    hash_set = set(hashes)
    for password in passwords:
        pw_hash = generate_ntlm_hash(password)
        print(f"[*] Testing password '{password}' (NT hash: {pw_hash})...")
        if pw_hash in hash_set:
            cracked[pw_hash] = password
            print(f"    [+] match: {pw_hash}:{password}")
    print(f"[+] Custom check: {len(cracked)} of {len(passwords)} password(s) matched")
    return cracked


def bulk_lookup(hash_type, hashes):
    """Look up hashes in chunks of 300. Returns {hash: password} for hits."""
    cracked = {}
    url = f"{API_BASE_URL}?hashtype={hash_type}"
    i = 0
    total = len(hashes)

    while i < total:
        chunk = hashes[i:i + 300]
        print(f"[*] Looking up {i + 1}-{min(i + 300, total)} of {total}...")
        try:
            resp = requests.post(
                url,
                headers={"Content-Type": "text/plain"},
                data="\n".join(chunk),
            )
            if resp.status_code == 200:
                for line in resp.text.strip().split("\n"):
                    if ':' in line and '[not found]' not in line.lower():
                        h, pw = line.split(':', 1)
                        h, pw = h.strip().lower(), pw.strip()
                        if h and pw:
                            cracked[h] = pw
                            print(f"    [+] {h}:{pw}")
                i += 300
                if i < total:
                    time.sleep(5)
            elif resp.status_code == 429:
                print("[!] Rate limit (429). Waiting 15 minutes before retrying...",
                      file=sys.stderr)
                time.sleep(900)  # retry same chunk, do not advance i
            else:
                print(f"[!] Unexpected status {resp.status_code}; skipping chunk",
                      file=sys.stderr)
                i += 300
                time.sleep(5)
        except requests.exceptions.RequestException as e:
            print(f"[!] Network error: {e}; skipping chunk", file=sys.stderr)
            i += 300

    print(f"[+] Cracked {len(cracked)} of {total} hashes")
    return cracked


# --------------------------------------------------------------------------- #
# 3. Report                                                                    #
# --------------------------------------------------------------------------- #
def parse_hash_file(file_path):
    """Parse every user:rid:lm:nt::: line into structured rows."""
    data = []
    with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            parts = line.split(':')
            if len(parts) >= 4:
                domain_user = parts[0]
                if '\\' in domain_user:
                    domain, username = domain_user.split('\\', 1)
                else:
                    domain, username = "", domain_user
                data.append({
                    'Domain': domain,
                    'Username': username,
                    'UID': parts[1],
                    'LM Hash': parts[2],
                    'NT Hash': parts[3],
                })
    return data


def match_passwords(hash_data, cracked):
    """Return domain/user/password rows for accounts whose hash was cracked."""
    matched = []
    for entry in hash_data:
        nt = entry['NT Hash'].lower()
        lm = entry['LM Hash'].lower()
        password = cracked.get(nt) or cracked.get(lm)
        if password:
            matched.append({
                'Domain': entry['Domain'],
                'Username': entry['Username'],
                'Password': password,
            })
    return matched


def style_sheet(ws, has_data=True):
    if not has_data:
        return
    header_font = Font(name='Titillium Web', size=11, bold=True)
    header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
    data_font = Font(name='Titillium Web', size=10)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = data_font
    for column in ws.columns:
        max_len = max((len(str(c.value)) for c in column if c.value is not None),
                      default=0)
        ws.column_dimensions[column[0].column_letter].width = min(max_len + 2, 50)


def save_to_excel(hash_data, matched, output_file):
    with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
        if hash_data:
            pd.DataFrame(hash_data).to_excel(writer, sheet_name='All_Hashes', index=False)
        else:
            pd.DataFrame([['No hash data found']], columns=['Message']).to_excel(
                writer, sheet_name='All_Hashes', index=False)
        if matched:
            pd.DataFrame(matched).to_excel(writer, sheet_name='Cracked_Passwords', index=False)
        else:
            pd.DataFrame([['No cracked passwords found']], columns=['Message']).to_excel(
                writer, sheet_name='Cracked_Passwords', index=False)

    wb = load_workbook(output_file)
    style_sheet(wb['All_Hashes'], has_data=bool(hash_data))
    style_sheet(wb['Cracked_Passwords'], has_data=bool(matched))
    wb.save(output_file)
    print(f"[+] Sheet 1 - All_Hashes: {len(hash_data)} entries")
    print(f"[+] Sheet 2 - Cracked_Passwords: {len(matched)} entries")
    print(f"[+] Excel report created: {output_file}")


# --------------------------------------------------------------------------- #
# 4. Export  (domain:username:password)                                        #
# --------------------------------------------------------------------------- #
def parse_cracked_text(file_path):
    """Read a hash:password file (e.g. ntlm.pw results) -> {hash: password}."""
    cracked = {}
    with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
        for line in f:
            line = line.strip()
            if not line or '[not found]' in line.lower() or ':' not in line:
                continue
            h, pw = line.split(':', 1)
            h, pw = h.strip().lower(), pw.strip()
            if h and pw:
                cracked[h] = pw
    return cracked


def rows_from_xlsx(xlsx_path):
    """Read domain/username/password from the report's cracked sheet."""
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        header = [str(c.value).strip().lower() if c.value else "" for c in next(ws.iter_rows())]
        if "username" in header and "password" in header:
            ui, pi = header.index("username"), header.index("password")
            di = header.index("domain") if "domain" in header else None
            rows = []
            for r in ws.iter_rows(min_row=2, values_only=True):
                if not r or ui >= len(r) or pi >= len(r) or not r[ui] or not r[pi]:
                    continue
                domain = (r[di] if di is not None and di < len(r) else "") or ""
                rows.append((str(domain).strip(), str(r[ui]).strip(), str(r[pi]).strip()))
            return rows
    print("[!] No sheet with Username + Password columns found.", file=sys.stderr)
    sys.exit(1)


def export_creds(input_file, dump_file, output_file):
    """Write domain:username:password lines."""
    rows = []
    if input_file.lower().endswith((".xlsx", ".xlsm")):
        for domain, user, pw in rows_from_xlsx(input_file):
            # username may already be DOMAIN/user
            u = user.replace("\\", "/")
            if not domain and "/" in u:
                domain, u = u.split("/", 1)
            rows.append((domain, u.split("/")[-1], pw))
    else:
        if not dump_file:
            print("[!] -f is a hash:password text file, so the dump is needed to "
                  "resolve domain/username.\n"
                  "    Run:  python3 lookhash.py <dump-file> -f results.txt -export -o file.txt",
                  file=sys.stderr)
            sys.exit(1)
        cracked = parse_cracked_text(input_file)
        for entry in parse_hash_file(dump_file):
            pw = cracked.get(entry['NT Hash'].lower()) or cracked.get(entry['LM Hash'].lower())
            if pw:
                rows.append((entry['Domain'], entry['Username'], pw))

    seen, out_lines = set(), []
    for domain, user, pw in rows:
        line = f"{domain}:{user}:{pw}"
        if line not in seen:
            seen.add(line)
            out_lines.append(line)

    with open(output_file, 'w', encoding='utf-8') as f:
        f.write("\n".join(out_lines) + ("\n" if out_lines else ""))
    print(f"[+] Exported {len(out_lines)} domain:username:password lines -> {output_file}")


# --------------------------------------------------------------------------- #
# main                                                                         #
# --------------------------------------------------------------------------- #
def main():
    parser = argparse.ArgumentParser(
        description="Extract NT hashes, look them up via ntlm.pw, and build an Excel report "
                    "(or export cracked creds as domain:username:password).")
    parser.add_argument("dump", nargs="?", help="Input secretsdump-style hash file "
                        "(required for lookup mode, and for -export from a hash:password file)")
    parser.add_argument("-f", "--file", help="Export mode input: a hash:password file "
                        "(e.g. results.txt) or a report .xlsx")
    parser.add_argument("-export", "--export", action="store_true",
                        help="Export domain:username:password to -o instead of building a report")
    parser.add_argument("-o", "--output", required=True, help="Output file (.xlsx report, or .txt for -export)")
    parser.add_argument("-t", "--type", default="nt",
                        choices=["nt", "lm", "md5", "sha1", "sha256"],
                        help="Hash type to look up (default: nt)")
    parser.add_argument("--no-lookup", action="store_true",
                        help="Skip the ntlm.pw lookup; build report from hashes only")
    parser.add_argument("-p", "--password", action="append", default=[], metavar="PASSWORD",
                        help="Candidate password to test locally against the hashes "
                             "(NTLM match). Repeat -p for several. Combines with the "
                             "ntlm.pw lookup unless --no-lookup is given.")
    args = parser.parse_args()

    if args.export:
        if not args.file:
            parser.error("-export requires -f (the cracked results file or report .xlsx)")
        export_creds(args.file, args.dump, args.output)
        return

    if not args.dump:
        parser.error("the dump file is required (positional) unless using -export")

    hashes = extract_nt_hashes(args.dump)
    cracked = {} if args.no_lookup else bulk_lookup(args.type, hashes)
    if args.password:
        cracked.update(custom_lookup(hashes, args.password))

    hash_data = parse_hash_file(args.dump)
    matched = match_passwords(hash_data, cracked)
    save_to_excel(hash_data, matched, args.output)


if __name__ == "__main__":
    main()
